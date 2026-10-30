"""웹문서 대량제작 상품 키워드 목록 생성 — python keywords/gen_webdoc_keywords.py

축 조합으로 키워드 후보를 만들고, 같은 의도는 하나의 매핑 URL로 묶는다(카니발 방지).
검색량은 없는 후보 목록이다 — 네이버 검색광고 키워드도구로 조회수를 붙여 하위 키워드를 걸러낼 것.
"""
import collections
import openpyxl
from openpyxl.styles import Font, PatternFill

# ---- 핵심 상품어: (키워드, 매핑 URL 슬러그, 묶음) ----
CORE = [
    ('웹문서 대량제작', '웹문서-대량제작', '웹문서'), ('웹문서 제작', '웹문서-대량제작', '웹문서'),
    ('웹문서 대량생성', '웹문서-대량제작', '웹문서'), ('웹문서 배포', '웹문서-배포', '웹문서'),
    ('웹문서 발행', '웹문서-배포', '웹문서'), ('웹문서 대량배포', '웹문서-배포', '웹문서'),
    ('웹문서 마케팅', '네이버-웹문서-노출', '웹문서'), ('웹문서 홍보', '네이버-웹문서-노출', '웹문서'),
    ('웹문서 광고', '네이버-웹문서-노출', '웹문서'), ('네이버 웹문서', '네이버-웹문서-노출', '웹문서'),
    ('네이버 웹문서 노출', '네이버-웹문서-노출', '웹문서'), ('네이버 웹문서 등록', '네이버-웹문서-노출', '웹문서'),
    ('구글 웹문서 노출', '네이버-웹문서-노출', '웹문서'), ('구글 검색 노출', '네이버-웹문서-노출', '웹문서'),
    ('웹페이지 대량제작', '웹문서-대량제작', '대량페이지'), ('웹페이지 대량생성', '웹문서-대량제작', '대량페이지'),
    ('홈페이지 대량제작', '웹문서-대량제작', '대량페이지'), ('랜딩페이지 대량제작', '웹문서-대량제작', '대량페이지'),
    ('대량 랜딩페이지', '웹문서-대량제작', '대량페이지'), ('키워드 페이지 제작', '웹문서-대량제작', '대량페이지'),
    ('키워드별 페이지 제작', '웹문서-대량제작', '대량페이지'), ('키워드 랜딩페이지', '웹문서-대량제작', '대량페이지'),
    ('지역 키워드 페이지', '지역-키워드-페이지', '대량페이지'), ('지역별 페이지 제작', '지역-키워드-페이지', '대량페이지'),
    ('지역 랜딩페이지 제작', '지역-키워드-페이지', '대량페이지'), ('롱테일 키워드 페이지', '지역-키워드-페이지', '대량페이지'),
    ('프로그래매틱 SEO', '프로그래매틱-SEO', 'SEO'), ('SEO 페이지 제작', '프로그래매틱-SEO', 'SEO'),
    ('SEO 웹문서', '프로그래매틱-SEO', 'SEO'), ('홈페이지 SEO', '프로그래매틱-SEO', 'SEO'),
    ('SEO 대행', '프로그래매틱-SEO', 'SEO'), ('검색엔진 최적화 대행', '프로그래매틱-SEO', 'SEO'),
    ('GEO 최적화', 'GEO-AEO-최적화', 'AI검색'), ('AEO 최적화', 'GEO-AEO-최적화', 'AI검색'),
    ('생성형 검색 최적화', 'GEO-AEO-최적화', 'AI검색'), ('AI 검색 최적화', 'GEO-AEO-최적화', 'AI검색'),
    ('AI 검색 노출', 'GEO-AEO-최적화', 'AI검색'), ('챗GPT 노출', 'GEO-AEO-최적화', 'AI검색'),
    ('네이버 AI 브리핑 노출', 'GEO-AEO-최적화', 'AI검색'), ('구글 AI 개요 노출', 'GEO-AEO-최적화', 'AI검색'),
    ('타지역 홍보', '타지역-홍보', '타지역'), ('타지역 노출', '타지역-홍보', '타지역'),
    ('타지역 마케팅', '타지역-홍보', '타지역'), ('출장 업체 홍보', '타지역-홍보', '타지역'),
    ('출장 업종 마케팅', '타지역-홍보', '타지역'), ('지역 확장 마케팅', '타지역-홍보', '타지역'),
    ('블로그 대체 마케팅', '블로그-대안-마케팅', '타지역'), ('블로그 대신 홍보', '블로그-대안-마케팅', '타지역'),
    ('사이트맵 제출 대행', '웹문서-배포', 'SEO'), ('서치어드바이저 등록 대행', '웹문서-배포', 'SEO'),
    ('서치콘솔 등록 대행', '웹문서-배포', 'SEO'), ('검색엔진 등록 대행', '웹문서-배포', 'SEO'),
]
# 구매 의도 수식어 — 같은 매핑 URL로 간다 (서비스 페이지가 받는다)
BUY_MODS = ['비용', '가격', '견적', '업체', '대행', '추천', '최저가', '단가', '잘하는곳', '사례']
# 정보 의도 수식어 — 블로그 글로 간다
INFO_MODS = ['방법', '효과', '뜻', '차이', '장단점', '주의사항']
INFO_OK = {'웹문서', 'SEO', 'AI검색', '대량페이지', '타지역'}

# ---- 업종군: 지역·키워드 검색으로 문의가 들어오는 업종 (사장님이 직접 홍보 수단을 찾는 업종) ----
# 금융·법률은 광고 규제 업종 — 납품 페이지는 대부업 등록번호·이자율 표기, 변호사 광고규정(결과 보장·허위 사례 금지)을 지킨다
# 업종·서비스는 사이트와 같은 데이터에서 읽는다 (data/webdoc/industries.json — 페이지와 키워드가 같은 원본)
import json as _json, os as _os
_DATA = _json.load(open(_os.path.join(_os.path.dirname(_os.path.abspath(__file__)), '..', 'data', 'webdoc', 'industries.json'), encoding='utf-8'))
INDUSTRY_GROUPS = {}
for _it in _DATA:
    INDUSTRY_GROUPS.setdefault(_it['group'], []).append(_it['name'])
SERVICES = []  # (서비스명, 업종명) — 서비스 페이지가 있는 것만 (사이트 규칙: note 있고 업종명과 다르고 slug 안 겹침)
_taken = {f"{n}-홍보" for n in (i['name'] for i in _DATA)}
for _it in _DATA:
    for _s in _it['services']:
        if isinstance(_s, dict) and _s.get('note') and _s['name'] != _it['name'] and f"{_s['name']}-홍보" not in _taken:
            _taken.add(f"{_s['name']}-홍보")
            SERVICES.append((_s['name'], _it['name']))
INDUSTRIES = [i for g in INDUSTRY_GROUPS.values() for i in g]
GROUP_OF = {i: g for g, lst in INDUSTRY_GROUPS.items() for i in lst}
IND_ANGLES = [  # (키워드 패턴, URL 패턴, 의도)
    ('{i} 웹문서 제작', '{i}-홍보', '구매'), ('{i} 웹문서 대량제작', '{i}-홍보', '구매'),
    ('{i} 지역 페이지 제작', '{i}-홍보', '구매'), ('{i} 홍보', '{i}-홍보', '구매'),
    ('{i} 업체 홍보', '{i}-홍보', '구매'), ('{i} 홍보 방법', '{i}-홍보', '정보'),
    ('{i} 마케팅', '{i}-홍보', '구매'), ('{i} 업체 광고', '{i}-홍보', '구매'),
    ('{i} 타지역 홍보', '{i}-홍보', '구매'), ('{i} 네이버 노출', '{i}-홍보', '정보'),
    ('{i} 업체 손님 늘리는 법', '{i}-홍보', '정보'), ('{i} 키워드 광고 대안', '{i}-홍보', '정보'),
    ('{i} 업체 홈페이지 제작', '{i}-홍보', '구매'), ('{i} 지역 키워드', '{i}-홍보', '구매'),
    ('{i} 파워링크 비용', '{i}-홍보', '정보'), ('{i} 플레이스 노출', '{i}-홍보', '정보'),
    ('{i} 광고 효과 없을 때', '{i}-홍보', '정보'),
]
# 지역 × 주력 업종 — "인천 하수구 업체 홍보"처럼 사장님이 자기 영업지역을 붙여 찾는 검색. 새 페이지 없이 업종 홍보 페이지로 묶는다
TOP_INDUSTRIES = ['하수구', '누수탐지', '이사', '입주청소', '크레인', '인테리어', '방수', '에어컨설치', '철거', '열쇠',
                  '방역', '용달', '도배', '보일러', '샷시', '커튼', '블라인드', '대출', '개인회생', '변호사', '법무사', '세무사']

# ---- 지역: 상품 구매자가 지역을 붙여 찾는 경우 (대행사·업체를 동네에서 찾는 검색) ----
REGIONS = [
    '서울', '부산', '대구', '인천', '광주', '대전', '울산', '세종', '경기', '강원', '충북', '충남', '전북', '전남',
    '경북', '경남', '제주', '수원', '성남', '용인', '고양', '일산', '부천', '안산', '안양', '화성', '동탄', '평택',
    '시흥', '김포', '파주', '의정부', '남양주', '하남', '광명', '군포', '오산', '이천', '양주', '구리', '포천',
    '청주', '천안', '아산', '전주', '창원', '김해', '포항', '구미', '원주', '춘천', '강남', '송파', '마포', '영등포',
]
REG_ANGLES = [
    ('{r} 웹문서 제작', '웹문서-대량제작', '구매'), ('{r} 웹문서 대량제작', '웹문서-대량제작', '구매'),
    ('{r} SEO 대행', '웹문서-대량제작', '구매'), ('{r} 지역 키워드 마케팅', '웹문서-대량제작', '구매'),
]

# ---- 정보성 주제 (블로그 자동 발행 큐) — 사장님이 문제를 느끼는 순간의 검색 ----
INFO_TOPICS = [
    '블로그 상위노출 안될 때', '블로그 저품질 확인 방법', '블로그 글 누락 해결', '블로그 체험단 효과 없을 때',
    '네이버 플레이스 거리 제한', '플레이스 타지역 노출 안되는 이유', '출장 업체 타지역 손님 받는 법',
    '영업 지역 넓은 업체 홍보 방법', '웹문서 탭이란', '네이버 웹문서 노출 원리', '웹문서와 블로그 차이',
    '웹문서와 홈페이지 차이', 'SEO GEO AEO 차이', 'AI 검색에 내 업체 나오게 하는 법', '챗GPT가 추천하는 업체 되는 법',
    '네이버 AI 브리핑에 노출되는 글', '사이트맵 제출 방법', '서치어드바이저 웹페이지 수집 요청', '구글 색인 요청 방법',
    '네이버 색인 안될 때', '중복 문서 판정 피하는 법', '양산형 페이지 패널티', '도어웨이 페이지란',
    '키워드 광고비 아끼는 법', '파워링크 대신 할 수 있는 홍보', '지역 키워드 찾는 법', '롱테일 키워드란',
    '검색량 적은 키워드가 돈 되는 이유', '키워드 하나에 페이지 하나', '페이지 수가 많으면 노출이 늘까',
    '웹문서 대량제작 업체 고르는 기준', '상위노출 보장 업체를 피해야 하는 이유', '웹문서 제작 후 노출까지 걸리는 기간',
    '홈페이지 있는데 검색이 안 될 때', '내 업체 이름 검색해도 안 나올 때', '네이버 구글 동시에 노출하는 법',
    '썸네일이 검색 클릭률에 미치는 영향', '구조화 데이터란', '캐러셀 노출 조건', '전화 문의 늘리는 페이지 구성',
]
INFO_VARS = ['', ' (2026년 기준)', ' — 출장 업체 편', ' 총정리', ' Q&A']


def build():
    rows, seen = [], set()

    def add(kw, url, cat, intent, use, ind='', reg=''):
        if kw in seen:
            return
        seen.add(kw)
        rows.append([kw, kw.replace(' ', ''), cat, intent, use, GROUP_OF.get(ind, ''), ind, reg, f'/웹문서/{url}/' if use == '페이지' else '/blog/'])

    for kw, url, grp in CORE:
        add(kw, url, f'핵심-{grp}', '구매', '페이지')
    for kw, url, grp in CORE:
        for m in BUY_MODS:
            add(f'{kw} {m}', url, f'핵심-{grp}', '구매', '페이지')
        if grp in INFO_OK:
            for m in INFO_MODS:
                add(f'{kw} {m}', url, f'핵심-{grp}', '정보', '블로그')
    for i in INDUSTRIES:
        for p, u, intent in IND_ANGLES:
            add(p.format(i=i), u.format(i=i), '업종', intent, '페이지' if intent == '구매' else '블로그', ind=i)
    for sv, ind in SERVICES:
        for p, intent in [('{s} 업체 홍보', '구매'), ('{s} 홍보', '구매'), ('{s} 마케팅', '구매'), ('{s} 웹문서', '구매'), ('{s} 업체 손님 늘리는 법', '정보')]:
            add(p.format(s=sv), f'{sv}-홍보', '서비스', intent, '페이지' if intent == '구매' else '블로그', ind=ind)
    for r in REGIONS:
        for i in TOP_INDUSTRIES:
            add(f'{r} {i} 업체 홍보', f'{i}-홍보', '지역×업종', '구매', '페이지', ind=i, reg=r)
    for r in REGIONS:
        for p, u, intent in REG_ANGLES:
            add(p.format(r=r), u.format(r=r), '지역', intent, '페이지', reg=r)
    for t in INFO_TOPICS:
        for v in INFO_VARS:
            add(t + v, '', '정보성', '정보', '블로그')
    return rows


def main():
    rows = build()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '전체 키워드'
    hdr = ['No', '키워드', '붙여쓰기', '분류', '검색 의도', '용도', '업종군', '업종', '지역', '매핑 URL']
    ws.append(hdr)
    for n, r in enumerate(rows, 1):
        ws.append([n] + r)

    pages = collections.OrderedDict()
    for r in rows:
        if r[4] == '페이지':
            pages.setdefault(r[8], []).append(r[0])
    ps = wb.create_sheet('페이지 목록')
    ps.append(['No', '매핑 URL', 'H1(대표 키워드)', '분류', '묶인 키워드 수', '묶인 키워드'])
    cat_of = {r[8]: r[2] for r in rows if r[4] == '페이지'}
    for n, (url, kws) in enumerate(pages.items(), 1):
        ps.append([n, url, kws[0], cat_of[url], len(kws), ', '.join(kws)])

    s = wb.create_sheet('요약', 0)
    by_cat = collections.Counter(r[2].split('-')[0] for r in rows)
    by_use = collections.Counter(r[4] for r in rows)
    s.append(['항목', '개수'])
    s.append(['전체 키워드', len(rows)])
    s.append(['페이지로 가는 키워드', by_use['페이지']])
    s.append(['블로그 큐로 가는 키워드', by_use['블로그']])
    s.append(['고유 페이지(매핑 URL)', len(pages)])
    s.append([])
    s.append(['분류', '키워드 수'])
    for k, v in by_cat.items():
        s.append([k, v])
    s.append([])
    s.append(['업종군', '업종 수', '키워드 수'])
    by_grp = collections.Counter(r[5] for r in rows if r[5])
    for g, lst in INDUSTRY_GROUPS.items():
        s.append([g, len(lst), by_grp[g]])
    s.append([])
    s.append(['주의', '검색량 미조회 후보 목록. 네이버 검색광고 키워드도구로 월간 조회수를 붙여 0~10회 키워드는 페이지 대신 블로그 큐로 내릴 것.'])

    for sh in wb.worksheets:
        for c in sh[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill('solid', fgColor='DDE7F5')
        sh.freeze_panes = 'A2'
        for col in sh.columns:
            sh.column_dimensions[col[0].column_letter].width = min(60, max(10, max(len(str(c.value or '')) for c in col) + 2))
    out = 'keywords/웹문서대량제작_키워드.xlsx'
    wb.save(out)
    print(f'키워드 {len(rows)} / 페이지 {len(pages)} / 블로그 {by_use["블로그"]} → {out}')
    print('분류:', dict(by_cat))


if __name__ == '__main__':
    main()
